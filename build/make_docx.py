# -*- coding: utf-8 -*-
"""组装需求分析文档 docx。

用法（在仓库根目录执行）：
    python build/make_docx.py
输出：docs/需求分析文档.docx
"""
import glob
import os
import struct
import sys

from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
IMG_DIR = os.path.join(ROOT, "uml")
OUT = os.path.join(ROOT, "docs", "需求分析文档.docx")

SONG = "宋体"
HEI = "黑体"

# 版面参数（A4，页边距上下 2.54cm、左右 3.17cm）
CONTENT_W = 14.66          # 纵向版心宽 cm
CONTENT_H = 24.62          # 纵向版心高 cm
LAND_W = 23.36             # 横向版心宽 cm
MAX_FIG_H = 13.0           # 纵向页单图最大高度 cm
LAND_MAX_FIG_H = 12.0      # 横向页单图最大高度 cm


def set_run_font(run, cn_font, size, bold=False, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn_font)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def add_field(paragraph, code, placeholder=""):
    run = paragraph.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = code
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = placeholder
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    run._r.append(f1)
    run._r.append(it)
    run._r.append(f2)
    run._r.append(t)
    run._r.append(f3)
    return run


def png_size(path):
    with open(path, "rb") as fh:
        head = fh.read(24)
    w, h = struct.unpack(">II", head[16:24])
    return w, h


def find_image(key):
    pattern = os.path.join(IMG_DIR, "图%s_*.png" % key)
    hits = sorted(glob.glob(pattern))
    if not hits:
        raise FileNotFoundError("找不到图 %s 的 PNG：%s" % (key, pattern))
    return hits[0]


def setup_section(sec, landscape=False):
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.5)


def add_header_footer(sec, title):
    header = sec.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, SONG, 9, color=RGBColor(0x59, 0x59, 0x59))
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("第 ")
    set_run_font(r1, SONG, 9)
    add_field(fp, "PAGE")
    r2 = fp.add_run(" 页")
    set_run_font(r2, SONG, 9)
    for r in fp.runs:
        set_run_font(r, SONG, 9)


def style_table(table, widths_cm, font_size=9):
    tbl = table._tbl
    tblPr = tbl.tblPr
    # 固定列宽
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths_cm):
                cell.width = Cm(widths_cm[idx])
    # 表头底纹与加粗
    hdr = table.rows[0]
    for cell in hdr.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "D9E2F3")
        tcPr.append(shd)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs:
                set_run_font(r, HEI, font_size, bold=True)
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    set_run_font(r, SONG, font_size)


def add_caption(doc, text, is_table=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, HEI, 10.5, bold=True)
    return p


def add_figure_block(doc, info, stats):
    key = info["key"]
    path = find_image(key)
    pw, ph = png_size(path)
    ar = ph / pw
    landscape = bool(info.get("landscape"))
    if landscape:
        # 新起一页（新节自带换页，不额外插分页符，避免出现空白页）
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        setup_section(sec, landscape=True)
        add_header_footer(sec, "基于本地 OCR 的 PDF 文件显示与管理系统 · 需求分析文档")
        w_cm = LAND_W
        if w_cm * ar > LAND_MAX_FIG_H:
            w_cm = LAND_MAX_FIG_H / ar
        target_w, target_h = w_cm, w_cm * ar
    else:
        w_cm = CONTENT_W
        if w_cm * ar > MAX_FIG_H:
            w_cm = MAX_FIG_H / ar
        target_w, target_h = w_cm, w_cm * ar

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(path, width=Cm(target_w))
    add_caption(doc, info["caption"], is_table=False)
    for line in info.get("note", []):
        np_ = doc.add_paragraph()
        np_.paragraph_format.first_line_indent = Pt(21)
        np_.paragraph_format.space_after = Pt(3)
        np_.paragraph_format.line_spacing = 1.3
        run = np_.add_run(line)
        set_run_font(run, SONG, 10.5)

    if landscape:
        # 回到纵向节
        doc.add_section(WD_SECTION.NEW_PAGE)
        sec2 = doc.sections[-1]
        setup_section(sec2, landscape=False)
        add_header_footer(sec2, "基于本地 OCR 的 PDF 文件显示与管理系统 · 需求分析文档")

    stats["figures"] += 1
    stats["fig_cm"] += target_h + 1.2


def add_table_block(doc, info, stats):
    add_caption(doc, info["caption"], is_table=True)
    rows = info["rows"]
    header = info["header"]
    table = doc.add_table(rows=len(rows) + 1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(header):
        cell = table.cell(0, i)
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, HEI, 9, bold=True)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, SONG, 9)
    style_table(table, info.get("widths", [CONTENT_W / len(header)] * len(header)))
    stats["tables"] += 1
    stats["table_rows"] += len(rows) + 1


def add_cover(doc, info):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(info["title"])
    set_run_font(run, HEI, 22, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(info["subtitle"])
    set_run_font(r2, HEI, 14)
    for _ in range(3):
        doc.add_paragraph()
    table = doc.add_table(rows=len(info["rows"]), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, (k, v) in enumerate(info["rows"]):
        c0 = table.cell(ri, 0)
        c0.text = ""
        r = c0.paragraphs[0].add_run(k)
        set_run_font(r, HEI, 11, bold=True)
        c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c1 = table.cell(ri, 1)
        c1.text = ""
        r = c1.paragraphs[0].add_run(v)
        set_run_font(r, SONG, 11)
    style_table(table, [3.4, 11.2], font_size=11)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)


def main():
    import doc_part1, doc_part2a, doc_part2b, doc_part3, doc_part4, doc_part5, doc_part6

    parts = [
        doc_part1.BLOCKS,
        doc_part2a.BLOCKS,
        doc_part2b.BLOCKS,
        doc_part3.BLOCKS,
        doc_part4.BLOCKS,
        doc_part5.BLOCKS,
        doc_part6.BLOCKS,
    ]

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), HEI)

    sec = doc.sections[0]
    setup_section(sec, landscape=False)
    sec.different_first_page_header_footer = True
    add_header_footer(sec, "基于本地 OCR 的 PDF 文件显示与管理系统 · 需求分析文档")

    stats = {"figures": 0, "tables": 0, "table_rows": 0, "fig_cm": 0.0, "chars": 0, "paras": 0}

    for blocks in parts:
        for blk in blocks:
            kind = blk[0]
            if kind == "cover":
                add_cover(doc, blk[1])
            elif kind == "pagebreak":
                doc.add_page_break()
            elif kind == "toc":
                p = doc.add_paragraph()
                add_field(p, r'TOC \o "1-2" \h \z \u', "（请在 Word 中右键选择“更新域”生成目录）")
                for r in p.runs:
                    set_run_font(r, SONG, 12)
            elif kind in ("h1", "h2", "h3"):
                level = {"h1": 1, "h2": 2, "h3": 3}[kind]
                p = doc.add_paragraph(style="Heading %d" % level)
                p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.3
                run = p.add_run(blk[1])
                set_run_font(run, HEI, {1: 16, 2: 14, 3: 12}[level], bold=True)
                stats["chars"] += len(blk[1])
                stats["paras"] += 1
            elif kind == "p":
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Pt(24)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.5
                run = p.add_run(blk[1])
                set_run_font(run, SONG, 12)
                stats["chars"] += len(blk[1])
                stats["paras"] += 1
            elif kind == "bullets":
                for item in blk[1]:
                    p = doc.add_paragraph(style="List Bullet")
                    p.paragraph_format.left_indent = Pt(24)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.3
                    run = p.add_run(item)
                    set_run_font(run, SONG, 12)
                    stats["chars"] += len(item)
                    stats["paras"] += 1
            elif kind == "table":
                add_table_block(doc, blk[1], stats)
            elif kind == "figure":
                add_figure_block(doc, blk[1], stats)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)

    # 粗略页数估算
    text_pages = stats["chars"] / 1250.0
    fig_pages = stats["figures"] * 0.62
    table_pages = stats["table_rows"] / 42.0
    total = 3 + text_pages + fig_pages + table_pages
    print("已生成：%s" % OUT)
    print("图片 %d 张，表格 %d 张（%d 行），正文字数约 %d 字"
          % (stats["figures"], stats["tables"], stats["table_rows"], stats["chars"]))
    print("估算页数：文字 %.1f + 图 %.1f + 表 %.1f + 前置 3 = %.1f 页"
          % (text_pages, fig_pages, table_pages, total))


if __name__ == "__main__":
    main()
