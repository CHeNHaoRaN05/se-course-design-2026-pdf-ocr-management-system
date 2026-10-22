"""Acceptance tests for a laboratory document workflow."""
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
import io
import threading

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select, update

from ocr_service.db import File, Job, TextIndex, now
from ocr_service.queue import claim, heartbeat, save_page
from ocr_service.tables import ruled_tables
from ocr_service.worker import Worker
from .conftest import authorization
from .documents import image_bytes, inventory_image, pdf_bytes

API = "/api/v1/ocr"


class InventoryEngine:
    """Deterministic service double; it does not evaluate OCR accuracy."""
    def __init__(self, unavailable=False):
        self.unavailable = unavailable
        self.calls = 0

    def recognize(self, image, options):
        self.calls += 1
        if self.unavailable:
            raise RuntimeError("simulated accelerator unavailable")
        return {"width": image.width, "height": image.height, "text": "显微镜借用登记", "corrected": False,
                "tables": [], "blocks": [{"id": 8, "text": "显微镜借用登记", "confidence": .72,
                "bbox": [.08, .08, .65, .16], "needs_review": True}]}


def upload(service, pdf=False):
    response = service.client.post(API + "/files", files={"file": ("inventory.pdf" if pdf else "inventory.png",
        pdf_bytes() if pdf else image_bytes(9))})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def enqueue(service, file_id, **kwargs):
    response = service.client.post(API + "/jobs", json={"file_id": file_id, **kwargs})
    assert response.status_code == 202, response.text
    return response.json()["id"]


def test_document_review_drives_search_and_download(service):
    document = upload(service, pdf=True)
    task = enqueue(service, document, pages=[3, 1])
    Worker(service.config, service.sessions, InventoryEngine()).run_once()
    response = service.client.get(f"{API}/jobs/{task}/pages").json()
    assert [p["page_no"] for p in response] == [1, 3]
    blocks = response[0]["result"]["blocks"]
    blocks[0].update(text="显微镜已归还", needs_review=False)
    revision = {"expected_revision": response[0]["revision"], "blocks": blocks}
    assert service.client.patch(f"{API}/jobs/{task}/pages/1", json=revision).json()["revision"] == 2
    assert service.client.patch(f"{API}/jobs/{task}/pages/1", json=revision).status_code == 409
    hits = service.client.get(API + "/search?q=归还").json()["items"]
    assert len(hits) == 1 and hits[0]["page_no"] == 1
    for extension in ["txt", "json", "csv", "html", "xlsx"]:
        downloaded = service.client.get(f"{API}/jobs/{task}/export?format={extension}")
        assert downloaded.status_code == 200 and downloaded.headers["x-ocr-exported-pages"] == "2"
    assert service.client.get(f"{API}/jobs/{task}/pages/3/image").status_code == 200


def test_private_inventory_is_not_visible_to_other_labs(service):
    document = upload(service)
    task = enqueue(service, document)
    other = authorization(service.config, "another-lab")
    for suffix in [f"/files/{document}/content", f"/jobs/{task}", f"/jobs/{task}/pages"]:
        assert service.client.get(API + suffix, headers=other).status_code == 404
    assert service.client.get(API + "/jobs", headers=other).json() == []
    assert service.client.get(API + "/search?q=登记", headers=other).json()["items"] == []
    readonly = authorization(service.config, role="viewer")
    assert service.client.post(f"{API}/jobs/{task}/cancel", headers=readonly).status_code == 403
    assert service.client.get(API + "/usage", headers={"Authorization": "Bearer invalid"}).status_code == 401


def test_resource_validation_and_optional_model_contract(service):
    document = upload(service)
    for body in [{"pages": [-1]}, {"pages": [2]}, {"mode": "handwriting"}, {"recognition_model": "custom"}]:
        assert service.client.post(API + "/jobs", json={"file_id": document, **body}).status_code == 422
    capability = service.client.get(API + "/capabilities").json()
    assert capability["models"] == ["official"] and capability["modes"] == ["general"]
    assert service.client.post(API + "/files", files={"file": ("invalid.pdf", b"%PDF-0 broken")}).status_code == 422
    service.config.max_upload_bytes = 16
    assert service.client.post(API + "/files", files={"file": ("inventory.png", image_bytes())}).status_code == 413


def test_unavailable_engine_can_recover_after_manual_retry(service):
    task = enqueue(service, upload(service))
    engine = InventoryEngine(unavailable=True)
    worker = Worker(service.config, service.sessions, engine)
    worker.run_once()
    assert service.client.get(f"{API}/jobs/{task}").json()["status"] == "retry_wait"
    worker.run_once()
    assert service.client.get(f"{API}/jobs/{task}").json()["status"] == "failed"
    engine.unavailable = False
    assert service.client.post(f"{API}/jobs/{task}/retry").status_code == 202
    worker.run_once()
    assert service.client.get(f"{API}/jobs/{task}").json()["status"] == "succeeded"


def test_racing_consumers_and_cancelled_lease(service):
    task = enqueue(service, upload(service))
    start = threading.Barrier(2)
    def consume():
        start.wait(timeout=5)
        return claim(service.sessions, service.config)
    with ThreadPoolExecutor(max_workers=2) as pool:
        tickets = list(pool.map(lambda _: consume(), range(2)))
    lease, = [t for t in tickets if t]
    assert lease[0] == task
    assert service.client.post(f"{API}/jobs/{task}/cancel").status_code == 200
    assert not save_page(service.sessions, *lease, 1, error="late_result")
    assert not heartbeat(service.sessions, service.config, *lease)


def test_dead_worker_is_replaced_without_accepting_stale_writes(service):
    task = enqueue(service, upload(service))
    abandoned = claim(service.sessions, service.config)
    with service.sessions.begin() as session:
        session.execute(update(Job).where(Job.id == task).values(lease_until=now() - timedelta(minutes=1)))
    replacement = claim(service.sessions, service.config)
    assert replacement[1] != abandoned[1]
    assert not save_page(service.sessions, *abandoned, 1, error="outdated")
    assert save_page(service.sessions, *replacement, 1, error="current")


def test_budget_reservation_failure_is_atomic(service):
    document = upload(service)
    def exhausted(*args):
        raise HTTPException(429, "Document quota exhausted")
    service.app.state.ocr_before_enqueue = exhausted
    assert service.client.post(API + "/jobs", json={"file_id": document}).status_code == 429
    with service.sessions() as session:
        assert session.get(File, document).active_job_id is None
        assert session.scalar(select(Job.id)) is None


def test_reviewed_spreadsheet_values_are_literal(service):
    task = enqueue(service, upload(service))
    Worker(service.config, service.sessions, InventoryEngine()).run_once()
    page = service.client.get(f"{API}/jobs/{task}/pages/1").json()
    blocks = page["result"]["blocks"]
    blocks[0]["text"] = '\x00=HYPERLINK("https://example.invalid")'
    assert service.client.patch(f"{API}/jobs/{task}/pages/1", json={"expected_revision": 1, "blocks": blocks}).status_code == 200
    exported = service.client.get(f"{API}/jobs/{task}/export?format=xlsx")
    book = load_workbook(io.BytesIO(exported.content))
    assert book["OCR文本"]["C2"].data_type == "s"
    assert book["OCR文本"]["C2"].value.startswith("'=")


def test_inventory_grid_coordinates():
    with inventory_image(11) as image:
        grids = ruled_tables(image, [{"text": "显微镜", "bbox": [70/960, 292/560, 220/960, 330/560]}])
    assert len(grids) == 1
    assert len(grids[0]["cells"]) == 3 and len(grids[0]["cells"][0]) == 2
    assert grids[0]["cells"][1][0] == "显微镜"


def test_new_generation_keeps_its_search_result(service):
    document = upload(service)
    earlier = enqueue(service, document)
    worker = Worker(service.config, service.sessions, InventoryEngine())
    worker.run_once()
    current = enqueue(service, document)
    worker.run_once()
    old_page = service.client.get(f"{API}/jobs/{earlier}/pages/1").json()
    assert service.client.patch(f"{API}/jobs/{earlier}/pages/1", json={"expected_revision": 1, "blocks": old_page["result"]["blocks"]}).status_code == 409
    with service.sessions() as session:
        assert session.scalar(select(TextIndex.job_id)) == current
